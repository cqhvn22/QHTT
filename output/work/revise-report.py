from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml.ns import qn
src=Path(r'C:/Users/chuqu/OneDrive/Desktop/QHTT.docx'); d=Document(src)
original_code=[p.text for p in d.paragraphs if p.style.name=='Ma nguon']
body=d.element.body; ps=d.paragraphs
start=body.index(ps[12]._p); end=body.index(ps[89]._p)
for el in list(body)[start:end]: body.remove(el)
anchor=ps[89]._p
# New paragraphs are moved before the existing source appendix.
def move(p):
 body.remove(p._p); anchor.addprevious(p._p); return p
def p(text='',style='Normal'):
 a=move(d.add_paragraph(text,style)); f=a.paragraph_format
 f.first_line_indent=Cm(1.27); f.space_after=Pt(6); f.space_before=Pt(0); f.line_spacing=1.15
 for r in a.runs:
  r.font.name='Times New Roman'; r.font.size=Pt(14); r.font.color.rgb=RGBColor(0,0,0)
 return a
def h(t):
 a=p(t,'Heading 1'); a.paragraph_format.keep_with_next=True
 for r in a.runs:r.bold=True
 return a
def sub(t):
 a=p(t,'Heading 2'); a.paragraph_format.keep_with_next=True
 for r in a.runs:r.bold=True
 return a
def page(title):
 a=h(title); a.paragraph_format.page_break_before=True
 return a
def pic(file,cap):
 a=p(); a.paragraph_format.keep_with_next=True
 a.add_run().add_picture(str(Path('output/work')/file),width=Cm(15.4)); p(cap,'Caption')
page('1 Giới thiệu bài toán')
p('Trong hoạt động sản xuất, vận tải và phân bổ nguồn lực, người quản lý cần lựa chọn một phương án đạt hiệu quả tốt nhất trong điều kiện nguồn lực có hạn. Khi hàm mục tiêu và các điều kiện ràng buộc đều là biểu thức tuyến tính, vấn đề có thể được mô hình hóa thành bài toán quy hoạch tuyến tính.')
sub('1.1 Mô hình toán học')
p('Bài toán gồm các biến quyết định x1, x2, …, xn; hàm mục tiêu cần tìm giá trị lớn nhất hoặc nhỏ nhất; hệ ràng buộc và điều kiện dấu của các biến.')
p('- Hàm mục tiêu: f(x) = c1x1 + c2x2 + … + cnxn → Max hoặc Min.')
p('- Ràng buộc thứ i: ai1x1 + ai2x2 + … + ainxn ≤, = hoặc ≥ bi. Mỗi ràng buộc chọn một dấu quan hệ phù hợp.')
p('- Miền biến: xj ≥ 0, xj ≤ 0 hoặc xj tự do, tùy theo ý nghĩa của biến trong bài toán.')
p('Miền nghiệm là tập các phương án thỏa mãn đồng thời mọi ràng buộc. Phương án tối ưu là phương án khả thi có giá trị hàm mục tiêu tốt nhất. Một bài toán cũng có thể vô nghiệm hoặc có hàm mục tiêu không bị chặn theo chiều tối ưu.')
sub('1.2 Mục tiêu của đề tài')
p('Đề tài xây dựng website QHTT để người học nhập mô hình, giải bài toán và theo dõi lời giải. Chương trình hỗ trợ phương pháp đơn hình và trực quan hóa bằng hình học cho bài toán hai biến, giúp đối chiếu kết quả tính toán với miền nghiệm trên đồ thị.')
p('- Đầu vào: chiều tối ưu, các hệ số, dấu ràng buộc và miền của biến.')
p('- Đầu ra: trạng thái bài toán, nghiệm và giá trị tối ưu nếu có, cùng các bước giải hoặc đồ thị minh họa.')
page('2 Yêu cầu của hệ thống')
sub('2.1 Yêu cầu chức năng')
p('- Nhập bài toán: tạo tiêu đề, chọn Max hoặc Min, số biến và số ràng buộc; nhập hệ số mục tiêu, ma trận ràng buộc, vế phải và miền biến.')
p('+ Cho phép dùng số nguyên, số thập phân và phân số; hỗ trợ dán dữ liệu dạng bảng từ Excel, thêm ràng buộc và xóa dòng cuối.')
p('- Kiểm tra dữ liệu: phát hiện giá trị không hợp lệ, mẫu số bằng 0 hoặc kích thước bài toán không phù hợp và hiển thị thông báo để người dùng sửa.')
p('- Giải bằng đơn hình: chuẩn hóa bài toán, xử lý biến giả, thực hiện đổi cơ sở và trình bày các bảng đơn hình cùng kết luận.')
p('- Giải bằng hình học: dựng đường biên, miền nghiệm, các đỉnh và đường mức của hàm mục tiêu cho bài toán hai biến có miền nghiệm bị chặn.')
p('+ Cho phép thu phóng, di chuyển đồ thị, xem thông tin điểm và điều khiển chuyển động đường mức.')
p('- Quản lý bài toán: mở nhiều tab, tạo bài mới, tạo bài mẫu, đóng tab và mở lại tab vừa đóng.')
p('- Lưu trữ và xuất kết quả: lưu đề dưới dạng JSON, mở lại JSON, tùy chọn tự động lưu phiên làm việc và in hoặc lưu lời giải thành PDF.')
p('- Điều khiển tính toán: giải bằng nút Giải hoặc Ctrl + Enter; cho phép dừng và giải lại sau khi thay đổi dữ liệu.')
sub('2.2 Yêu cầu phi chức năng và phạm vi')
p('- Giao diện bằng tiếng Việt, trình bày rõ vùng nhập dữ liệu và vùng kết quả; sử dụng trực tiếp trên trình duyệt và có thể chạy ngoại tuyến.')
p('- Tác vụ tính toán chạy trong Web Worker để giảm ảnh hưởng đến thao tác trên giao diện; có giới hạn thời gian và số lần đổi cơ sở.')
p('- Phạm vi nhập liệu: từ 1 đến 20 biến và tối đa 40 ràng buộc. Chương trình xử lý biến liên tục, không giải bài toán nguyên, nhị phân hoặc phi tuyến.')
p('- Dữ liệu phiên làm việc nằm trên trình duyệt của người dùng; không có chức năng tài khoản hoặc đồng bộ qua máy chủ.')
page('3 Công nghệ sử dụng')
p('Sản phẩm được xây dựng bằng HTML, CSS và JavaScript thuần. Mã nguồn chạy ở phía trình duyệt, không cần cài đặt máy chủ ứng dụng hoặc hệ quản trị cơ sở dữ liệu.')
p('- HTML5: tạo cấu trúc trang, biểu mẫu nhập hệ số, các nút thao tác, vùng lời giải và hộp thoại hướng dẫn. Tệp chính là index.html.')
p('- CSS3: định dạng bố cục, bảng dữ liệu, nút bấm, vùng kết quả và bản in. Các quy tắc giao diện được đặt trong css/style.css.')
p('- JavaScript: xử lý dữ liệu nhập, sự kiện giao diện, thuật toán giải, báo cáo kết quả và quản lý các tab bài toán.')
p('- BigInt và lớp số hữu tỉ: biểu diễn tử số, mẫu số bằng số nguyên để thực hiện phép tính phân số trong bộ giải, hạn chế sai số do biểu diễn số thực.')
p('- Web Worker và Blob: tạo tác vụ tính toán nền từ mã JavaScript. Cách tổ chức này cho phép chạy bộ giải khi mở index.html trực tiếp từ máy tính.')
p('- SVG: biểu diễn đường biên, miền nghiệm, các điểm và đường mức trong phần hình học. Hình vẽ hỗ trợ phóng to và tương tác với người dùng.')
p('- localStorage: lưu dữ liệu phiên làm việc khi bật Tự động lưu. Dữ liệu gắn với trình duyệt và địa chỉ mở ứng dụng.')
p('- JSON và các API tệp của trình duyệt: đóng gói đề bài để tải xuống, đọc lại và chuyển bài toán giữa các máy.')
p('- Chức năng in của trình duyệt: in lời giải hoặc lưu thành PDF thông qua nút In / lưu PDF.')
p('- Phông Computer Modern nhúng trong math-fonts.js: phục vụ hiển thị biểu thức toán học khi chạy ngoại tuyến. Tệp chứa dữ liệu phông; chương trình không phụ thuộc dịch vụ phông trực tuyến.')
sub('3.1 Môi trường thực hiện')
p('Người sử dụng cần một trình duyệt hỗ trợ JavaScript, BigInt, Web Worker và SVG. Trong lần chạy minh họa của báo cáo, ứng dụng được mở trực tiếp bằng Microsoft Edge và cho kết quả ở cả phương pháp đơn hình lẫn hình học.')
page('4 Thiết kế và nguyên lý hoạt động')
sub('4.1 Tổ chức mã nguồn')
p('- index.html và css/style.css: xây dựng cấu trúc và định dạng giao diện.')
p('- js/app.js: điều phối nhập liệu, thao tác tab, khởi chạy bộ giải, lưu mở tệp và cập nhật kết quả.')
p('- js/solver.js: phép tính số hữu tỉ, chuẩn hóa mô hình và thuật toán đơn hình.')
p('- js/geometry.js và js/geometry-view.js: tính toán miền nghiệm hai biến, dựng lời giải hình học và đồ thị tương tác.')
p('- js/report.js: tạo nội dung lời giải và nội dung phục vụ in.')
p('- js/session.js: đóng gói và khôi phục phiên làm việc; js/examples.js: các bài toán mẫu.')
p('- js/layout.js: điều chỉnh vùng nhập liệu; js/math-fonts.js: dữ liệu phông chữ toán học nhúng.')
sub('4.2 Luồng xử lý')
p('Người dùng nhập dữ liệu → chương trình kiểm tra dữ liệu → chuyển bài toán vào bộ giải trong Web Worker → trả kết quả → giao diện hiển thị tóm tắt và lời giải → người dùng lưu đề hoặc in kết quả.')
sub('4.3 Phương pháp đơn hình')
p('- Chuẩn hóa miền biến và ràng buộc: đổi dấu biến không dương, tách biến tự do, xử lý vế phải âm, thêm biến phụ hoặc biến dư; bổ sung biến giả khi cần tạo cơ sở ban đầu.')
p('- Với phương pháp Đơn hình trên giao diện, biến giả được xử lý bằng M tượng trưng. Bộ giải tính các ước lượng, chọn biến vào và biến ra rồi thực hiện phép đổi cơ sở.')
p('- Lặp đến khi thỏa điều kiện tối ưu hoặc phát hiện vô nghiệm, hàm mục tiêu không bị chặn hay chạm giới hạn tính toán. Bộ giải có giới hạn 45 giây và 600 phép đổi cơ sở.')
sub('4.4 Phương pháp hình học')
p('Với hai biến, chương trình xác định các đường biên, tìm giao điểm khả thi và kiểm tra miền nghiệm. Khi miền bị chặn, chương trình tính giá trị mục tiêu ở các đỉnh và thể hiện đường mức để minh họa vị trí tối ưu. Miền không bị chặn được thông báo là ngoài phạm vi lời giải hình học của giao diện.')
page('5 Cài đặt và khởi chạy')
p('- Bước 1: Sao chép hoặc giải nén thư mục DonHinhWeb. Giữ index.html cùng cấp với hai thư mục css và js.')
p('- Bước 2: Nhấp chuột phải vào index.html, chọn Open with rồi chọn Microsoft Edge hoặc Google Chrome.')
p('- Bước 3: Kiểm tra giao diện có vùng Bài toán và Lời giải. Không cần cài thêm thư viện, cơ sở dữ liệu hoặc máy chủ để mở chương trình.')
pic('01-giao-dien.png','Hình 1 Giao diện chương trình khi khởi chạy')
p('- Vùng bên trái dùng để chọn phương pháp và nhập bài. Vùng bên phải trình bày kết quả. Thanh phía trên hỗ trợ mở bài, lưu bài và tự động lưu.')
p('- Nhấn dấu + hoặc Bài mới để tạo tab trống. Tạo bài mẫu nạp một đề mẫu; sau đó cần nhấn Giải để tính.')
page('6 Chạy thử với một bài toán cụ thể')
sub('6.1 Phát biểu và mô hình hóa')
p('Xét một cơ sở sản xuất hai loại sản phẩm. Gọi x1 và x2 là lượng sản phẩm cần sản xuất, được xem là các đại lượng liên tục. Lợi ích trên mỗi đơn vị lần lượt là 3 và 2. Tổng lượng sản phẩm không vượt quá 4; riêng loại thứ nhất không vượt quá 2 và loại thứ hai không vượt quá 3. Cần tìm phương án đem lại lợi ích lớn nhất.')
p('- Hàm mục tiêu: Max f = 3x1 + 2x2.')
p('- Các ràng buộc: x1 + x2 ≤ 4; x1 ≤ 2; x2 ≤ 3.')
p('- Điều kiện dấu: x1 ≥ 0 và x2 ≥ 0.')
sub('6.2 Nhập dữ liệu')
p('- Bước 1: Chọn Đơn hình, đặt tiêu đề Ví dụ hướng dẫn.')
p('- Bước 2: Nhập Số biến là 2, Ràng buộc là 3 rồi nhấn Áp dụng.')
p('- Bước 3: Chọn Max; nhập hệ số mục tiêu 3 và 2; đặt miền cả hai biến là ≥ 0.')
p('- Bước 4: Nhập từng ràng buộc theo thứ tự hệ số x1, hệ số x2, dấu, vế phải.')
p('+ Dòng 1: 1, 1, ≤, 4.')
p('+ Dòng 2: 1, 0, ≤, 2.')
p('+ Dòng 3: 0, 1, ≤, 3.')
p('- Bước 5: Nhấn Giải hoặc Ctrl + Enter và đợi kết quả.')
p('Nhập 0 cho hệ số không xuất hiện trong đề. Có thể nhập số thập phân như 0.5 hoặc phân số như 3/7. Khi thay đổi dữ liệu, cần nhấn Giải lại để cập nhật lời giải.')
page('7 Kết quả chạy bằng phương pháp đơn hình')
pic('02-don-hinh.png','Hình 2 Kết quả chạy bài toán bằng phương pháp đơn hình')
p('Kết quả thực tế: nghiệm tối ưu duy nhất là x1 = 2, x2 = 2; giá trị tối ưu f = 10. Chương trình hiển thị 2 lần đổi cơ sở. Người dùng cuộn vùng lời giải để xem quá trình chuẩn hóa và các bảng đơn hình.')
sub('7.1 Đối chiếu kết quả')
p('- Tính khả thi: 2 + 2 = 4 ≤ 4; 2 ≤ 2; 2 ≤ 3; cả hai biến đều không âm.')
p('- Giá trị mục tiêu: f = 3 × 2 + 2 × 2 = 10.')
p('- Với mọi phương án khả thi, f = 2(x1 + x2) + x1 ≤ 2 × 4 + 2 = 10. Phương án (2; 2) đạt cận trên này nên là phương án tối ưu.')
p('- Ý nghĩa: trong mô hình minh họa, sản xuất 2 đơn vị mỗi loại cho lợi ích lớn nhất là 10 đơn vị giá trị.')
page('8 Kết quả chạy bằng phương pháp hình học')
p('Giữ nguyên đề bài, đổi Phương pháp thành Hình học (2 biến), nhấn Giải và cuộn đến phần đồ thị.')
pic('04-do-thi.png','Hình 3 Miền nghiệm và đường mức của bài toán hai biến')
p('Miền nghiệm bị chặn có các đỉnh (0; 0), (2; 0), (2; 2), (1; 3) và (0; 3). Giá trị mục tiêu tương ứng là 0, 6, 10, 9 và 6. Giá trị lớn nhất đạt tại (2; 2), trùng với kết quả đơn hình.')
p('- Vùng xanh thể hiện miền nghiệm; vùng đỏ thể hiện phần ngoài miền nghiệm. Đường mức dịch chuyển để minh họa sự thay đổi giá trị hàm mục tiêu.')
p('- Cuộn chuột để thu phóng, kéo để di chuyển, bấm điểm hoặc đường thẳng để xem thông tin. Vừa khung đặt lại vùng nhìn.')
p('- Tạm dừng và thanh a dùng để điều khiển đường mức. Nút Đường mức tối ưu đưa đường mức đến vị trí tối ưu.')
page('9 Lưu trữ kết quả và xử lý tình huống')
sub('9.1 Lưu mở bài toán và in lời giải')
p('- Nhấn Lưu để tải đề của tab hiện tại dưới dạng JSON. Tệp lưu các hệ số, miền biến và phương pháp; không thay thế bản in lời giải.')
p('- Nhấn Mở bài và chọn JSON đã lưu. Chương trình mở bài ở tab mới; nhấn Giải để tính lại.')
p('- Bật Tự động lưu để lưu dữ liệu các tab trên trình duyệt hiện tại. Khi mở lại bằng cùng trình duyệt và địa chỉ, đề được khôi phục; lời giải cần tính lại.')
p('+ Tắt Tự động lưu sẽ xóa bản lưu tự động. Để chuyển máy hoặc có bản sao riêng, sử dụng tệp JSON.')
p('- Sau khi giải, nhấn In / lưu PDF. Trong hộp thoại in, chọn máy in hoặc Save as PDF rồi xác nhận.')
p('- Dấu × đóng tab; nút ↶ mở lại tab vừa đóng. Chương trình giữ tối đa 10 tab đã đóng để mở lại.')
sub('9.2 Các tình huống cần chú ý')
p('- Dữ liệu không hợp lệ: kiểm tra hệ số, dấu ràng buộc, số biến, số ràng buộc và mẫu số của phân số.')
p('- Bài toán vô nghiệm: không tồn tại phương án thỏa tất cả điều kiện; cần đối chiếu các ràng buộc đã nhập với đề bài.')
p('- Hàm mục tiêu không bị chặn: có phương án khả thi nhưng không có giá trị tối ưu hữu hạn theo chiều đang xét.')
p('- Chạm giới hạn tính toán: chưa thể kết luận trạng thái toán học của bài toán; kiểm tra dữ liệu hoặc giảm quy mô mô hình.')
p('- Hình học yêu cầu đúng 2 biến và miền bị chặn. Khi không đáp ứng phạm vi này, chuyển sang Đơn hình nếu phù hợp.')
p('- Trang mất bố cục hoặc nút không chạy: kiểm tra đủ tệp css, js và JavaScript đã được cho phép. Nếu không khởi tạo được Web Worker, thử mở lại bằng trình duyệt hỗ trợ tính năng này.')
page('10 Đánh giá và kết luận')
sub('10.1 Kết quả đạt được')
p('Chương trình tổ chức được quy trình từ nhập mô hình đến tính toán và trình bày lời giải trên trình duyệt. Người dùng có thể lựa chọn đơn hình hoặc hình học, lưu đề để sử dụng lại và in kết quả phục vụ học tập.')
p('Trong bài chạy minh họa, hai phương pháp đều trả về nghiệm duy nhất (2; 2) và giá trị tối ưu 10. Kết quả được đối chiếu bằng cách thay vào các ràng buộc và chứng minh cận trên của hàm mục tiêu. Đây là kiểm tra cho ví dụ trình bày trong báo cáo, không phải kết luận về mọi bộ dữ liệu đầu vào.')
sub('10.2 Hạn chế')
p('- Chỉ xử lý mô hình tuyến tính với biến liên tục; chưa hỗ trợ quy hoạch nguyên, nhị phân hoặc phi tuyến.')
p('- Phần hình học chỉ phục vụ bài toán hai biến có miền nghiệm bị chặn.')
p('- Các giới hạn quy mô, thời gian và số phép đổi cơ sở có thể làm một số bài toán dừng trước khi có kết luận.')
p('- Tự động lưu chỉ hoạt động trong phạm vi trình duyệt và địa chỉ ứng dụng; chưa có tài khoản hoặc đồng bộ dữ liệu giữa thiết bị.')
sub('10.3 Hướng phát triển')
p('- Bổ sung bộ kiểm thử cho nhiều trường hợp: vô nghiệm, không bị chặn, suy biến, nhiều nghiệm tối ưu và các loại miền biến.')
p('- Mở rộng khả năng minh họa miền không bị chặn và cải thiện thao tác nhập các mô hình lớn.')
p('- Xem xét bổ sung chức năng giải bài toán nguyên và quản lý bài toán theo danh mục trong các phiên bản tiếp theo.')
p('Đề tài đáp ứng mục tiêu xây dựng công cụ hỗ trợ học quy hoạch tuyến tính: nhập mô hình, giải bài, theo dõi quá trình và đối chiếu kết quả bằng đồ thị trong phạm vi hai biến.')
page('Phụ lục Toàn bộ mã nguồn sản phẩm')
p('Phần này lưu toàn bộ mã nguồn có trong báo cáo, theo từng tệp HTML, CSS và JavaScript. Các dòng dài được Word tự xuống dòng khi hiển thị. Tệp math-fonts.js bao gồm dữ liệu phông chữ nhúng.')
for name in ['index.html','css/style.css','js/app.js','js/examples.js','js/geometry-view.js','js/geometry.js','js/layout.js','js/math-fonts.js','js/report.js','js/session.js','js/solver.js']:p('- '+name)
# Existing appendix begins on a fresh page, with its original source kept intact.
ps[89].paragraph_format.page_break_before=True
for name in ['Heading 1','Heading 2','Caption']:
 st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(14); st.font.color.rgb=RGBColor(0,0,0)
 for e in list(st.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
out=Path('output/QHTT_da_chinh_sua.docx');d.save(out)
check=Document(out); assert [p.text for p in check.paragraphs if p.style.name=='Ma nguon']==original_code
assert [[c.text for c in r.cells] for r in check.tables[0].rows]==[[c.text for c in r.cells] for r in Document(src).tables[0].rows]
print('Saved; preserved all',len(original_code),'source lines and cover member table')
